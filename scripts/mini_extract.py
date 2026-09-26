"""Tiny extractor (txt/csv/tsv/xlsx/pdf/images) following API.md §7.3 chunking.

Used by scripts/smoke_detect.py and eval/run_eval.py. Not the agent's extractor.
OCR is delegated to `ocr_fn(image_bytes, mime) -> OcrResult` (HTTP /ocr or llm.ocr).
"""
from __future__ import annotations

import csv
import hashlib
import io
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Optional

from detect_core.contracts import (
    Chunk, FileMeta, OcrResult, chunk_id_for, classify_folder, file_type_for,
)

OcrFn = Callable[[bytes, str], OcrResult]
TEXT_TYPES = {"txt", "md", "log", "json", "env", "ini", "yaml", "yml", "xml", "eml"}
IMAGE_TYPES = {"png", "jpg", "jpeg", "tiff"}
CHUNK_CHARS, OVERLAP = 4000, 150
ROWS_PER_BLOCK, MAX_ROWS, MAX_OCR_PAGES = 50, 5000, 10


def split_text(text: str) -> list[str]:
    if len(text) <= CHUNK_CHARS:
        return [text] if text.strip() else []
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + CHUNK_CHARS])
        i += CHUNK_CHARS - OVERLAP
    return out


def _column_chunks(rows: list[list[str]], page: int) -> list[dict]:
    if not rows:
        return []
    header, data = rows[0], rows[1:MAX_ROWS + 1]
    out = []
    for b in range(0, len(data), ROWS_PER_BLOCK):
        block = data[b:b + ROWS_PER_BLOCK]
        for ci, h in enumerate(header):
            vals = [(r[ci] if ci < len(r) and r[ci] is not None else "") for r in block]
            text = "\n".join(str(v) for v in vals)
            if text.strip():
                out.append({"text": text, "page": page, "row": 2 + b, "column": ci + 1,
                            "column_header": str(h) if h is not None else ""})
    return out


def _png(img) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def extract(path: Path, ocr_fn: Optional[OcrFn] = None, display_path: Optional[str] = None
            ) -> tuple[FileMeta, list[Chunk]]:
    raw = path.read_bytes()
    fhash = hashlib.sha256(raw).hexdigest()
    ftype = file_type_for(path.name)
    shown = display_path or str(path)
    meta = dict(file_path=shown, file_hash=fhash, file_type=ftype,
                folder_class=classify_folder(shown), size_bytes=len(raw),
                modified_at=datetime.fromtimestamp(os.path.getmtime(path), timezone.utc)
                .strftime("%Y-%m-%dT%H:%M:%SZ"))
    pieces: list[dict] = []
    reason = None
    try:
        if ftype in TEXT_TYPES:
            pieces = [{"text": t} for t in split_text(raw.decode("utf-8", errors="replace"))]
        elif ftype in ("csv", "tsv"):
            rows = list(csv.reader(io.StringIO(raw.decode("utf-8", errors="replace")),
                                   delimiter="\t" if ftype == "tsv" else ","))
            pieces = _column_chunks(rows, 1)
        elif ftype == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            for si, ws in enumerate(wb.worksheets, start=1):
                rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
                pieces.extend(_column_chunks(rows, si))
        elif ftype == "pdf":
            import pdfplumber
            ocr_pages = 0
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                for pno, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    if len(text.strip()) >= 20:
                        pieces.extend({"text": t, "page": pno} for t in split_text(text))
                    elif ocr_fn and ocr_pages < MAX_OCR_PAGES:
                        ocr_pages += 1
                        img = page.to_image(resolution=150).original
                        r = ocr_fn(_png(img), "image/png")
                        pieces.extend({"text": t, "page": pno, "ocr_confidence": r.confidence}
                                      for t in split_text(r.text))
                    elif ocr_fn:
                        reason = "ocr_page_cap"
        elif ftype in IMAGE_TYPES:
            if ocr_fn is None:
                return FileMeta(**meta, status="unscannable", status_reason="ocr_unavailable"), []
            mime = "image/jpeg" if ftype in ("jpg", "jpeg") else "image/png"
            data = raw
            if ftype == "tiff":
                from PIL import Image
                data = _png(Image.open(io.BytesIO(raw)))
            r = ocr_fn(data, mime)
            pieces = [{"text": t, "page": 1, "ocr_confidence": r.confidence} for t in split_text(r.text)]
        else:
            return FileMeta(**meta, status="unscannable", status_reason="unsupported_type"), []
    except Exception as exc:  # noqa: BLE001
        return FileMeta(**meta, status="unscannable", status_reason=f"extract_failed:{type(exc).__name__}"), []

    chunks = [Chunk(chunk_id=chunk_id_for(fhash, i), file_path=shown, file_hash=fhash,
                    file_type=ftype, folder_class=meta["folder_class"], page=p.get("page"),
                    row=p.get("row"), column=p.get("column"), column_header=p.get("column_header"),
                    ocr_confidence=p.get("ocr_confidence"), text=p["text"])
              for i, p in enumerate(pieces)]
    return FileMeta(**meta, status="ok", status_reason=reason), chunks
